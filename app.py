import os, math, io, time
from datetime import datetime, timezone
import requests
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from db_adapter import conn, init_db, engine
from baseline_filter import historical_name_match, BASELINE_SOURCE_ROWS, BASELINE_UNIQUE_NAMES

app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY','change-me')
DB_PATH = os.getenv('DB_PATH','places.db')
API_KEY = os.getenv('GOOGLE_MAPS_API_KEY','')
PLACES_URL = 'https://places.googleapis.com/v1/places:searchNearby'
MONTHLY_REQUEST_BUDGET = int(os.getenv('MONTHLY_REQUEST_BUDGET','4500'))

TYPE_GROUPS = {
 'Food & hospitality':['restaurant','cafe','bar','bakery','coffee_shop','pizza_restaurant','hotel'],
 'Beauty & wellness':['beauty_salon','hair_salon','barber_shop','nail_salon','spa','gym','fitness_center'],
 'Retail':['clothing_store','shoe_store','jewelry_store','electronics_store','furniture_store','hardware_store','florist','pet_store','book_store','supermarket'],
 'Automotive':['car_dealer','car_repair','car_rental','car_wash','tire_shop','auto_parts_store'],
 'Professional & local services':['accounting','lawyer','real_estate_agency','insurance_agency','travel_agency','laundry','dry_cleaning','locksmith','electrician','plumber','general_contractor','photographer'],
 'Health':['dentist','doctor','physiotherapist','pharmacy','medical_center','veterinary_care'],
 'Education':['school','preschool','driving_school','dance_school','music_school']
}
ALL_TYPES = sorted({x for g in TYPE_GROUPS.values() for x in g})
PRESETS = {
 'Milano':{'lat':45.4642,'lon':9.1900,'coverage_km':15.0},
 'Milano centro':{'lat':45.4642,'lon':9.1900,'coverage_km':5.0},
 'Monza':{'lat':45.5845,'lon':9.2744,'coverage_km':8.0},
 'Varese':{'lat':45.8206,'lon':8.8251,'coverage_km':8.0}
}

def clean_phone(v):
    if not v:return ''
    s=''.join(ch for ch in v if ch.isdigit() or ch=='+')
    if s.startswith('+39'): s=s[3:]
    elif s.startswith('0039'): s=s[4:]
    return s

def opening_date(v):
    if not v:return ''
    y=v.get('year');m=v.get('month');d=v.get('day')
    if not y:return ''
    return f'{y:04d}-{m:02d}-{d:02d}' if m and d else (f'{y:04d}-{m:02d}' if m else str(y))

def dest(lat,lon,meters,bearing):
    R=6371000.; br=math.radians(bearing); a=math.radians(lat); o=math.radians(lon); dr=meters/R
    a2=math.asin(math.sin(a)*math.cos(dr)+math.cos(a)*math.sin(dr)*math.cos(br))
    o2=o+math.atan2(math.sin(br)*math.sin(dr)*math.cos(a), math.cos(dr)-math.sin(a)*math.sin(a2))
    return math.degrees(a2),math.degrees(o2)

def grid(lat,lon,coverage_km,radius_m):
    spacing=max(250.,radius_m*1.35); rings=math.ceil(coverage_km*1000/spacing); out=[(lat,lon)]
    for r in range(1,rings+1):
        rr=r*spacing; n=max(6,math.ceil(2*math.pi*rr/spacing))
        for i in range(n): out.append(dest(lat,lon,rr,i*360/n))
    return out

def batches(xs,n=8):
    for i in range(0,len(xs),n): yield xs[i:i+n]

def nearby(lat,lon,radius,types,contacts=True):
    if not API_KEY: raise RuntimeError('GOOGLE_MAPS_API_KEY non configurata')
    fields=['places.id','places.displayName','places.formattedAddress','places.primaryType','places.types','places.businessStatus','places.openingDate','places.location','places.googleMapsUri']
    if contacts: fields += ['places.nationalPhoneNumber','places.websiteUri']
    payload={'includedTypes':types,'maxResultCount':20,'includeFutureOpeningBusinesses':True,
             'locationRestriction':{'circle':{'center':{'latitude':lat,'longitude':lon},'radius':float(radius)}}}
    h={'Content-Type':'application/json','X-Goog-Api-Key':API_KEY,'X-Goog-FieldMask':','.join(fields)}
    r=requests.post(PLACES_URL,json=payload,headers=h,timeout=30)
    if not r.ok: raise RuntimeError(f'Google Places API {r.status_code}: {r.text[:500]}')
    return r.json().get('places',[])

def norm(p):
    loc=p.get('location') or {}; name=(p.get('displayName') or {}).get('text','')
    return {'place_id':p.get('id',''),'name':name,'address':p.get('formattedAddress',''),'primary_type':p.get('primaryType',''),
            'types':','.join(p.get('types') or []),'business_status':p.get('businessStatus',''),'opening_date':opening_date(p.get('openingDate')),
            'phone':clean_phone(p.get('nationalPhoneNumber','')),'website':p.get('websiteUri',''),'maps_uri':p.get('googleMapsUri',''),
            'lat':loc.get('latitude'),'lon':loc.get('longitude')}

def monthly_requests_used(c, now_iso):
    month = now_iso[:7]
    row = c.execute("SELECT COALESCE(SUM(requests_count),0) AS n FROM scans WHERE substr(created_at,1,7)=?", (month,)).fetchone()
    return int(row['n'] or 0)

def run_scan(lat,lon,coverage_km,radius,selected_types,contacts=True):
    init_db(); c=conn(); baseline=c.execute('SELECT COUNT(*) n FROM places').fetchone()['n']==0; now=datetime.now(timezone.utc).isoformat(); used=monthly_requests_used(c,now); remaining=max(0,MONTHLY_REQUEST_BUDGET-used)
    if remaining <= 0:
        c.close(); raise RuntimeError(f'Budget mensile Google Places raggiunto ({MONTHLY_REQUEST_BUDGET} richieste).')
    cur=c.execute('INSERT INTO scans(created_at,center_lat,center_lon,coverage_km,cell_radius_m,types_count,is_baseline) VALUES(?,?,?,?,?,?,?)',(now,lat,lon,coverage_km,radius,len(selected_types),int(baseline)))
    scan_id=cur.lastrowid; c.commit(); seen=set(); new=set(); future=set(); req=0
    stop=False
    for glat,glon in grid(lat,lon,coverage_km,radius):
        if stop: break
        for b in batches(selected_types,8):
            if req >= remaining:
                stop=True; break
            try:
                places=nearby(glat,glon,radius,b,False); req+=1
            except RuntimeError as e:
                if '400' not in str(e): raise
                places=[]; tmp={}
                for t in b:
                    if req >= remaining:
                        stop=True; break
                    try:
                        part=nearby(glat,glon,radius,[t],False); req+=1
                        for x in part:
                            if x.get('id'): tmp[x['id']]=x
                    except RuntimeError as e2:
                        if '400' in str(e2): continue
                        raise
                places=list(tmp.values())
            for raw in places:
                p=norm(raw); pid=p['place_id']
                if not pid: continue
                seen.add(pid)
                if p['business_status']=='FUTURE_OPENING': future.add(pid)
                exists=c.execute('SELECT 1 FROM places WHERE place_id=?',(pid,)).fetchone()
                if not exists:
                    known = baseline or historical_name_match(p['name'])
                    if not known: new.add(pid)
                    c.execute('''INSERT INTO places VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(pid,p['name'],p['address'],p['primary_type'],p['types'],p['business_status'],p['opening_date'],p['phone'],p['website'],p['maps_uri'],p['lat'],p['lon'],now,now,scan_id,int(known)))
                else:
                    c.execute('''UPDATE places SET name=?,address=?,primary_type=?,types=?,business_status=?,opening_date=?,phone=CASE WHEN ?<>'' THEN ? ELSE phone END,website=CASE WHEN ?<>'' THEN ? ELSE website END,maps_uri=?,lat=?,lon=?,last_seen=? WHERE place_id=?''',(p['name'],p['address'],p['primary_type'],p['types'],p['business_status'],p['opening_date'],p['phone'],p['phone'],p['website'],p['website'],p['maps_uri'],p['lat'],p['lon'],now,pid))
            c.commit(); time.sleep(.01)
    c.execute('UPDATE scans SET requests_count=?,places_seen=?,new_places=?,future_openings=? WHERE id=?',(req,len(seen),len(new),len(future),scan_id)); c.commit(); c.close(); return scan_id,baseline

def get_scan(scan_id):
    c=conn(); s=c.execute('SELECT * FROM scans WHERE id=?',(scan_id,)).fetchone()
    if not s: c.close(); return None,[],[]
    new=[] if s['is_baseline'] else c.execute('SELECT * FROM places WHERE first_scan_id=? AND baseline=0 ORDER BY first_seen DESC',(scan_id,)).fetchall()
    fut=c.execute("SELECT * FROM places WHERE business_status='FUTURE_OPENING' ORDER BY opening_date,first_seen DESC").fetchall(); c.close(); return s,new,fut

@app.route('/status.md')
def status_md():
    init_db(); c=conn(); latest=c.execute('SELECT * FROM scans ORDER BY id DESC LIMIT 1').fetchone(); totals={'places':c.execute('SELECT COUNT(*) n FROM places').fetchone()['n'],'scans':c.execute('SELECT COUNT(*) n FROM scans').fetchone()['n'],'future':c.execute("SELECT COUNT(*) n FROM places WHERE business_status='FUTURE_OPENING'").fetchone()['n']}; now=datetime.now(timezone.utc).isoformat(); used=monthly_requests_used(c,now); c.close()
    lines=['# Google Maps New Openings — stato automatico','',f'Aggiornato automaticamente: {now}',f'Scansioni totali: {totals["scans"]}',f'Attività nel bacino: {totals["places"]}',f'Aperture future note: {totals["future"]}',f'Richieste Google Places usate nel mese: {used}/{MONTHLY_REQUEST_BUDGET}','']
    if latest:
        lines += ['## Ultima scansione',f'ID: {latest["id"]}',f'Data: {latest["created_at"]}',f'Copertura: {latest["coverage_km"]} km',f'Categorie interrogate: {latest["types_count"]}',f'Attività viste: {latest["places_seen"] or 0}',f'Nuove attività: {latest["new_places"] or 0}',f'Aperture future: {latest["future_openings"] or 0}',f'Richieste API: {latest["requests_count"] or 0}',f'Baseline: {"sì" if latest["is_baseline"] else "no"}']
    else:
        lines += ['## Ultima scansione','Nessuna scansione disponibile.']
    return Response('\n'.join(lines)+'\n',mimetype='text/markdown; charset=utf-8')

@app.route('/')
def index():
    init_db(); c=conn(); totals={'places':c.execute('SELECT COUNT(*) n FROM places').fetchone()['n'],'scans':c.execute('SELECT COUNT(*) n FROM scans').fetchone()['n'],'future':c.execute("SELECT COUNT(*) n FROM places WHERE business_status='FUTURE_OPENING'").fetchone()['n']}; latest=c.execute('SELECT * FROM scans ORDER BY id DESC LIMIT 10').fetchall(); c.close(); return render_template('index.html',presets=PRESETS,groups=TYPE_GROUPS,totals=totals,latest=latest,db_backend=engine.dialect.name,baseline_rows=BASELINE_SOURCE_ROWS,baseline_names=BASELINE_UNIQUE_NAMES)

@app.post('/scan')
def scan():
    try:
        name=request.form.get('preset','Milano'); p=PRESETS.get(name,PRESETS['Milano']); lat=float(request.form.get('lat') or p['lat']); lon=float(request.form.get('lon') or p['lon']); km=float(request.form.get('coverage_km') or p['coverage_km']); radius=float(request.form.get('cell_radius_m') or 7500); contacts=False; types=request.form.getlist('types') or ALL_TYPES
        sid,base=run_scan(lat,lon,km,radius,types,contacts); flash('Baseline creata: dalle prossime scansioni vedrai solo i nuovi Place ID.' if base else 'Scansione completata.','success'); return redirect(url_for('results',scan_id=sid))
    except Exception as e: flash(str(e),'danger'); return redirect(url_for('index'))

@app.route('/results/<int:scan_id>')
def results(scan_id):
    s,n,f=get_scan(scan_id); return render_template('results.html',scan=s,new_rows=n,future_rows=f) if s else redirect(url_for('index'))

@app.route('/export/<kind>/<int:scan_id>.xlsx')
def export_xlsx(kind,scan_id):
    s,n,f=get_scan(scan_id); rows=n if kind=='new' else f; wb=Workbook(); ws=wb.active; ws.title='Nuove attività' if kind=='new' else 'Aperture future'; hdr=['Ragione sociale','Numero','Categoria','Indirizzo','Stato','Data apertura','Sito','Google Maps','Place ID','Prima rilevazione']; ws.append(hdr)
    for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='D9EAF7'); c.alignment=Alignment(horizontal='center')
    for r in rows:
        ws.append([r['name'],str(r['phone'] or ''),r['primary_type'],r['address'],r['business_status'],r['opening_date'],r['website'],r['maps_uri'],r['place_id'],r['first_seen']]); ws.cell(ws.max_row,2).number_format='@'
    for i,w in enumerate([32,18,24,48,20,16,35,35,34,28],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2'; buf=io.BytesIO(); wb.save(buf); buf.seek(0); return send_file(buf,as_attachment=True,download_name=f'{kind}_scan_{scan_id}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

try:
    from autopilot import start_scheduler
    _scheduler=start_scheduler(__import__(__name__))
except Exception as e:
    print('Autopilot init failed:',e,flush=True)

if __name__=='__main__':
    init_db(); app.run(host='0.0.0.0',port=int(os.getenv('PORT','10000')),debug=False)
